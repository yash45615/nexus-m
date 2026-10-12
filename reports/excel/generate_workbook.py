from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[2]

PROCESSED_DIR = (
    PROJECT_ROOT
    / "data"
    / "processed"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "reports"
    / "excel"
    / "output"
)

OUTPUT_FILE = (
    OUTPUT_DIR
    / "NEXUS-M_Analyst_Model.xlsx"
)


# ---------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="5B9BD5",
)

SUBHEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)

GOOD_FILL = PatternFill(
    "solid",
    fgColor="C6EFCE",
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor="FFEB9C",
)

CRITICAL_FILL = PatternFill(
    "solid",
    fgColor="FFC7CE",
)


def style_title(
    ws,
    title: str,
) -> None:

    ws["A1"] = title

    ws["A1"].font = Font(
        bold=True,
        size=18,
        color="FFFFFF",
    )

    ws["A1"].fill = TITLE_FILL

    ws["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=8,
    )

    ws.row_dimensions[1].height = 28


def style_headers(ws) -> None:

    for cell in ws[1]:

        if cell.value is not None:

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = HEADER_FILL

            cell.alignment = Alignment(
                horizontal="center"
            )


def autofit_columns(ws) -> None:

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is None:
                continue

            value = str(cell.value)

            max_length = max(
                max_length,
                len(value),
            )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40,
        )


def add_table(
    ws,
    name: str,
) -> None:

    if ws.max_row < 2:
        return

    ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    table = Table(
        displayName=name,
        ref=ref,
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style

    ws.add_table(table)


# ---------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------

def load_csv(
    relative_path: str,
) -> pd.DataFrame:

    path = (
        PROCESSED_DIR
        / relative_path
    )

    if not path.exists():

        return pd.DataFrame()

    return pd.read_csv(path)


def load_json(
    relative_path: str,
) -> dict[str, Any]:

    path = (
        PROCESSED_DIR
        / relative_path
    )

    if not path.exists():

        return {}

    return json.loads(
        path.read_text(
            encoding="utf-8"
        )
    )


# ---------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------

def write_dataframe(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    start_col: int = 1,
) -> None:

    if df.empty:
        return

    for offset, column in enumerate(
        df.columns
    ):

        cell = ws.cell(
            row=start_row,
            column=start_col + offset,
        )

        cell.value = column

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = HEADER_FILL

    for row_offset, row in enumerate(
        df.itertuples(index=False),
        start=1,
    ):

        for col_offset, value in enumerate(
            row
        ):

            cell = ws.cell(
                row=start_row + row_offset,
                column=start_col + col_offset,
            )

            if pd.isna(value):
                cell.value = None
            else:
                cell.value = value


def add_metric(
    ws,
    row: int,
    label: str,
    value: Any,
) -> None:

    ws.cell(
        row=row,
        column=1,
        value=label,
    )

    ws.cell(
        row=row,
        column=1,
    ).font = Font(
        bold=True
    )

    ws.cell(
        row=row,
        column=2,
        value=value,
    )


# ---------------------------------------------------------------------
# Raw / analytical sheets
# ---------------------------------------------------------------------

def create_partner_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Partner Analysis"
    )

    style_title(
        ws,
        "NEXUS-M Partner Intelligence",
    )

    df = load_csv(
        "phase4/partner_performance.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Partner performance output "
            "is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


def create_region_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Region Analysis"
    )

    style_title(
        ws,
        "NEXUS-M Regional Performance",
    )

    df = load_csv(
        "phase4/regional_performance.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Regional performance output "
            "is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


def create_merchant_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Merchant Analysis"
    )

    style_title(
        ws,
        "NEXUS-M Merchant Intelligence",
    )

    df = load_csv(
        "phase8/merchant_segments.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Merchant segmentation output "
            "is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


def create_anomaly_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Anomaly Monitor"
    )

    style_title(
        ws,
        "NEXUS-M Anomaly Monitoring",
    )

    df = load_csv(
        "phase9/anomaly_results.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Anomaly output is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


def create_opportunity_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Opportunities"
    )

    style_title(
        ws,
        "NEXUS-M Opportunity Intelligence",
    )

    df = load_csv(
        "phase11/opportunity_priorities.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Opportunity output is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


def create_reconciliation_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Reconciliation"
    )

    style_title(
        ws,
        "NEXUS-M Reconciliation Analysis",
    )

    df = load_csv(
        "phase12/reconciliation_details.csv"
    )

    if df.empty:
        ws["A3"] = (
            "Reconciliation output "
            "is not available."
        )
        return

    write_dataframe(
        ws,
        df,
        start_row=3,
    )

    autofit_columns(ws)


# ---------------------------------------------------------------------
# Executive summary
# ---------------------------------------------------------------------


def create_executive_summary(
    wb: Workbook,
) -> None:

    ws = wb.active
    ws.title = "Executive Summary"

    style_title(
        ws,
        "NEXUS-M Executive Analytical Summary",
    )

    kpis = load_json("phase7/kpi_summary.json")

    if not kpis:
        kpis = load_json("phase4/core_kpis.json")

    metrics = (
        kpis.get("metrics", kpis)
        if isinstance(kpis, dict)
        else {}
    )

    # KPI table headers
    ws["A3"] = "Executive KPI"
    ws["B3"] = "Value"
    ws["C3"] = "Status"

    for cell in ws[3][:3]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = HEADER_FILL

    metric_names = [
        "transaction_count",
        "total_transaction_value",
        "success_rate",
        "failure_rate",
        "average_transaction_value",
        "average_latency_ms",
        "p95_latency_ms",
        "p99_latency_ms",
    ]

    row = 4

    for metric in metric_names:
        raw_value = metrics.get(metric)

        status = None

        # Handle structured KPI objects such as:
        # {"value": 251250.0, "status": "NO_THRESHOLD"}
        if isinstance(raw_value, dict):
            value = raw_value.get("value")
            status = raw_value.get("status")
        else:
            value = raw_value

        # Avoid writing unsupported objects into Excel.
        if isinstance(value, (dict, list, tuple, set)):
            value = json.dumps(value, default=str)

        ws.cell(
            row=row,
            column=1,
            value=metric,
        )

        ws.cell(
            row=row,
            column=2,
            value=value,
        )

        ws.cell(
            row=row,
            column=3,
            value=status,
        )

        row += 1

    # Business questions
    ws["E3"] = "Business Questions"
    ws["E3"].font = Font(
        bold=True,
        color="FFFFFF",
    )
    ws["E3"].fill = HEADER_FILL

    questions = [
        "How is the network performing?",
        "Where are failures concentrated?",
        "Which partners are underperforming?",
        "Which merchants are at risk?",
        "What are the major root causes?",
        "How much value is at risk?",
        "Which opportunities should be prioritized?",
        "Are settlements reconciling?",
    ]

    for index, question in enumerate(
        questions,
        start=4,
    ):
        ws.cell(
            row=index,
            column=5,
            value=question,
        )

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["E"].width = 60

    # Make the KPI values easier to read.
    for row_index in range(4, 4 + len(metric_names)):
        ws.cell(
            row=row_index,
            column=2,
        ).number_format = "#,##0.00"

# ---------------------------------------------------------------------
# Scenario analysis
# ---------------------------------------------------------------------

def create_scenario_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Scenario Analysis"
    )

    style_title(
        ws,
        "NEXUS-M Recovery Scenario Model",
    )

    ws["A3"] = "Input"
    ws["B3"] = "Value"

    for cell in [
        ws["A3"],
        ws["B3"],
    ]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = HEADER_FILL

    ws["A4"] = "Current Failed Value"
    ws["B4"] = 0

    ws["A5"] = "Current Failure Rate"
    ws["B5"] = 0

    ws["A6"] = "Target Failure Rate"
    ws["B6"] = 0.02

    ws["A8"] = "Scenario"
    ws["B8"] = "Recovery Rate"
    ws["C8"] = "Estimated Recovered Value"

    for cell in ws[8]:

        if cell.value is not None:

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = HEADER_FILL

    scenarios = [
        ("Low", 0.40),
        ("Base", 0.60),
        ("High", 0.80),
    ]

    row = 9

    for name, recovery_rate in scenarios:

        ws.cell(
            row=row,
            column=1,
            value=name,
        )

        ws.cell(
            row=row,
            column=2,
            value=recovery_rate,
        )

        ws.cell(
            row=row,
            column=3,
            value=f"=$B$4*B{row}",
        )

        row += 1

    ws["A14"] = (
        "Interpretation"
    )

    ws["A14"].font = Font(
        bold=True
    )

    ws["A15"] = (
        "Scenario outputs are analytical estimates "
        "and should not be interpreted as guaranteed "
        "financial outcomes."
    )

    ws.merge_cells(
        "A15:F17"
    )

    ws["A15"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 30


# ---------------------------------------------------------------------
# Lookup / configuration
# ---------------------------------------------------------------------

def create_lookup_sheet(
    wb: Workbook,
) -> None:

    ws = wb.create_sheet(
        "Lookup Config"
    )

    style_title(
        ws,
        "NEXUS-M Analytical Configuration",
    )

    rows = [
        ("Metric", "Threshold", "Interpretation"),
        ("Success Rate", 0.95, "Healthy above threshold"),
        ("Failure Rate", 0.05, "Warning above threshold"),
        ("P95 Latency", 500, "Warning above threshold"),
        ("P99 Latency", 1000, "Warning above threshold"),
        ("Opportunity Priority", 75, "High priority"),
        ("Reconciliation Score", 99, "Healthy"),
    ]

    for row_index, row in enumerate(
        rows,
        start=3,
    ):

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            if row_index == 3:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                )

                cell.fill = HEADER_FILL

    autofit_columns(ws)


# ---------------------------------------------------------------------
# Workbook creation
# ---------------------------------------------------------------------

def build_workbook() -> Path:

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = Workbook()

    create_executive_summary(wb)
    create_partner_sheet(wb)
    create_region_sheet(wb)
    create_merchant_sheet(wb)
    create_anomaly_sheet(wb)
    create_opportunity_sheet(wb)
    create_reconciliation_sheet(wb)
    create_scenario_sheet(wb)
    create_lookup_sheet(wb)

    # Freeze panes
    for ws in wb.worksheets:

        if ws.max_row > 2:
            ws.freeze_panes = "A3"

    wb.save(
        OUTPUT_FILE
    )

    return OUTPUT_FILE


if __name__ == "__main__":

    path = build_workbook()

    print(
        f"Excel workbook created: {path}"
    )
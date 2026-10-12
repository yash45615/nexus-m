from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "data" / "processed"
OUTPUT = ROOT / "reports" / "powerpoint" / "output"
OUTPUT_FILE = OUTPUT / "NEXUS-M_Executive_Briefing.pptx"

NAVY = RGBColor(23, 49, 79)
BLUE = RGBColor(43, 105, 170)
DARK = RGBColor(45, 52, 60)
GREY = RGBColor(105, 115, 125)


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except (OSError, pd.errors.ParserError, UnicodeDecodeError):
        return pd.DataFrame()


def first_existing(*paths: Path) -> Path | None:
    return next((path for path in paths if path.exists()), None)


def format_number(value: Any) -> str:
    try:
        number = float(value)
        if number.is_integer():
            return f"{int(number):,}"
        return f"{number:,.2f}"
    except (TypeError, ValueError):
        return str(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    size: int = 18,
    color=DARK,
    bold: bool = False,
):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str = ""):
    add_text(slide, title, 0.55, 0.35, 12.0, 0.55, 25, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.6, 0.95, 12.0, 0.4, 11, GREY)


def add_bullet_list(slide, items: list[str], x=0.8, y=1.65, width=11.7):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(5.2)
    )
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()

    for index, item in enumerate(items):
        paragraph = (
            frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        )
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(17)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(17)


def add_table(slide, df: pd.DataFrame, x, y, width, height, max_rows=7):
    if df.empty:
        add_text(
            slide,
            "No data available. Run the relevant analytics phase first.",
            x, y, width, height, 15, GREY,
        )
        return

    data = df.head(max_rows).copy()
    data = data.fillna("").astype(str)
    columns = list(data.columns[:5])

    table = slide.shapes.add_table(
        len(data) + 1,
        len(columns),
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    ).table

    for col_index, name in enumerate(columns):
        cell = table.cell(0, col_index)
        cell.text = name.replace("_", " ").title()
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY

        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(10)
            paragraph.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, (_, row) in enumerate(data.iterrows(), start=1):
        for col_index, name in enumerate(columns):
            cell = table.cell(row_index, col_index)
            cell.text = str(row[name])

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.color.rgb = DARK


def get_kpis() -> dict[str, Any]:
    path = first_existing(
        PROCESSED / "phase7" / "kpi_summary.json",
        PROCESSED / "phase4" / "core_kpis.json",
    )
    return read_json(path) if path else {}


def find_value(data: dict[str, Any], candidates: list[str]):
    for candidate in candidates:
        if candidate in data and not isinstance(data[candidate], (dict, list)):
            return data[candidate]

    for value in data.values():
        if isinstance(value, dict):
            result = find_value(value, candidates)
            if result is not None:
                return result

    return None


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: Cover
    slide = prs.slides.add_slide(blank)
    add_text(
        slide,
        "NEXUS-M",
        0.8, 1.55, 11.5, 0.9, 36, NAVY, True,
    )
    add_text(
        slide,
        "EXECUTIVE NETWORK INTELLIGENCE BRIEFING",
        0.85, 2.55, 11.5, 0.7, 23, BLUE, True,
    )
    add_text(
        slide,
        "Payment performance | Root-cause insights | Opportunities | Reconciliation",
        0.9, 3.45, 11.5, 0.65, 15, GREY,
    )
    add_text(
        slide,
        "Portfolio demonstration using synthetic data",
        0.9, 6.55, 11.5, 0.35, 11, GREY,
    )

    # Slide 2: Executive summary
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Executive Summary")
    kpis = get_kpis()

    count = find_value(
        kpis, ["transaction_count", "total_transactions", "transactions"]
    )
    success = find_value(kpis, ["success_rate"])
    failure = find_value(kpis, ["failure_rate"])
    value = find_value(
        kpis, ["transaction_value", "total_transaction_value"]
    )

    summary = [
        f"Transaction count: {format_number(count) if count is not None else 'Not available'}",
        f"Transaction value: {format_number(value) if value is not None else 'Not available'}",
        f"Success rate: {format_number(success) if success is not None else 'Not available'}",
        f"Failure rate: {format_number(failure) if failure is not None else 'Not available'}",
        "Review the partner, regional, and root-cause outputs before selecting operational priorities.",
        "Opportunity values are scenario estimates, not guaranteed financial outcomes.",
    ]
    add_bullet_list(slide, summary)

    # Slide 3: Partner performance
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Partner Performance", "Based on available Phase 4 outputs")
    partner_path = first_existing(
        PROCESSED / "phase4" / "partner_performance.csv",
        PROCESSED / "phase4" / "partner_key_kpis.csv",
    )
    partners = read_csv(partner_path) if partner_path else pd.DataFrame()
    add_table(slide, partners, 0.65, 1.55, 12.0, 4.9)

    # Slide 4: Regional performance
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Regional Performance")
    region_path = first_existing(
        PROCESSED / "phase4" / "regional_performance.csv",
        PROCESSED / "phase7" / "region_key_kpis.csv",
    )
    regions = read_csv(region_path) if region_path else pd.DataFrame()
    add_table(slide, regions, 0.65, 1.55, 12.0, 4.9)

    # Slide 5: Root-cause analysis
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Root-Cause Investigation")
    root_path = first_existing(
        PROCESSED / "phase10" / "partner_key_root_causes.csv",
        PROCESSED / "phase10" / "region_key_root_causes.csv",
    )
    root_causes = read_csv(root_path) if root_path else pd.DataFrame()
    add_table(slide, root_causes, 0.65, 1.55, 12.0, 4.9)

    # Slide 6: Opportunities
    slide = prs.slides.add_slide(blank)
    add_title(
        slide,
        "Opportunity Prioritization",
        "Estimated recovery scenarios; validate assumptions before action",
    )
    opportunities = read_csv(
        PROCESSED / "phase11" / "opportunity_priorities.csv"
    )
    add_table(slide, opportunities, 0.65, 1.55, 12.0, 4.9)

    # Slide 7: Reconciliation
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Settlement Reconciliation")
    recon_summary = read_json(
        PROCESSED / "phase12" / "reconciliation_summary.json"
    )
    recon_details = read_csv(
        PROCESSED / "phase12" / "exception_summary.csv"
    )

    if recon_summary:
        add_bullet_list(
            slide,
            [
                f"{key.replace('_', ' ').title()}: {format_number(value)}"
                for key, value in recon_summary.items()
                if not isinstance(value, (dict, list))
            ][:8],
            y=1.4,
        )
    else:
        add_text(
            slide,
            "Reconciliation summary is not available. Run Phase 12 first.",
            0.8, 1.6, 11.5, 0.6, 16, GREY,
        )

    if not recon_details.empty:
        add_table(slide, recon_details, 0.7, 4.3, 11.8, 2.1, max_rows=4)

    # Slide 8: Recommendations
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Recommended Next Actions")
    add_bullet_list(
        slide,
        [
            "Validate the most significant partner and regional performance changes against the underlying records.",
            "Investigate high-impact root-cause candidates and confirm whether they are operationally actionable.",
            "Prioritize opportunities using estimated recoverable value, evidence quality, and implementation effort.",
            "Resolve high-priority reconciliation exceptions and distinguish transaction exposure from actual discrepancies.",
            "Maintain data-quality checks, metric definitions, lineage, and reviewable analytical outputs.",
        ],
    )

    # Footer on every slide
    for number, slide in enumerate(prs.slides, start=1):
        add_text(
            slide,
            f"NEXUS-M | Synthetic-data portfolio project | {number}",
            0.65, 7.12, 12.0, 0.2, 8, GREY,
        )

    prs.save(OUTPUT_FILE)
    print(f"Presentation created: {OUTPUT_FILE}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()